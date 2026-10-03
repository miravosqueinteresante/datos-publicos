#!/usr/bin/env python3
"""Escaneo de datos personales en los XLSX crudos de reclamos.

Verifica si los archivos entregados contienen indicios de información sensible:
teléfonos, correos, cédulas o nombres propios en campos de texto libre.

No garantiza anonimización absoluta (eso solo lo puede certificar el emisor),
pero detecta patrones evidentes de PII y resume qué hay en el texto libre.

Uso:
    python scripts/scan_pii.py --fuente "C:\\...\\Reclamos Municipalidad"
"""

import argparse
import os
import re
import sys
from collections import Counter

import openpyxl

PHONE = re.compile(r"(?:\+?595\s?)?0?[234]\d{1,2}[\s.\-]?\d{3}[\s.\-]?\d{3}")
EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
CEDULA_KW = re.compile(r"\bc[ée]dula\w*\b|\bC\.?I\.?\s*(?:N[°º]?\.?)?\s*\d{3,8}\b", re.IGNORECASE)
DIGITS = re.compile(r"\b\d{7,9}\b")
NOMBRE_PROPIO = re.compile(r"\b[A-ZÁÉÍÓÚÑ][a-záéíóúñ]{2,}\s[A-ZÁÉÍÓÚÑ][a-záéíóúñ]{2,}\b")

FILES = [
    "1790948115_1_AO2023.xlsx",
    "1790948115_2_AO2024.xlsx",
    "1790948115_3_AO2025.xlsx",
    "1790948115_4_AO2026.xlsx",
]


def cells(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for c in row:
                if isinstance(c, str) and c.strip():
                    yield c.strip()
    wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fuente", default=r"C:\Users\pc\Desktop\Proyectos\Reclamos Municipalidad")
    args = parser.parse_args()

    total_phones = Counter()
    total_emails = Counter()
    total_cedula = Counter()
    total_digits = Counter()

    for fname in FILES:
        path = os.path.join(args.fuente, fname)
        if not os.path.exists(path):
            print("  WARN: no encontrado %s" % path)
            continue
        phones, emails, cedula, digits = Counter(), Counter(), Counter(), Counter()
        for c in cells(path):
            if PHONE.search(c):
                phones[PHONE.search(c).group(0)] += 1
            if EMAIL.search(c):
                emails[EMAIL.search(c).group(0)] += 1
            if CEDULA_KW.search(c):
                cedula[c[:80]] += 1
            elif DIGITS.search(c):
                # solo números largos sueltos (posibles cédulas/teléfonos)
                if len(DIGITS.findall(c)) == 1 and not re.search(r"20\d\d", c):
                    digits[c[:80]] += 1
        print("\n== %s ==" % fname)
        print("  teléfonos: %d | emails: %d | cédula/C.I.: %d | num. largos: %d"
              % (sum(phones.values()), sum(emails.values()),
                 sum(cedula.values()), sum(digits.values())))
        if cedula:
            for k, v in cedula.most_common(5):
                print("    [cédula?] %dx  %s" % (v, k))
        if digits:
            for k, v in digits.most_common(5):
                print("    [num largo] %dx  %s" % (v, k))

    # Nombre propio en texto libre (funcionarios, inspectores)
    print("\n== Nombres propios en texto libre (posibles funcionarios) ==")
    # escaneo de celdas largas (detalle/trazabilidad) en los 4 archivos
    for fname in FILES:
        path = os.path.join(args.fuente, fname)
        if not os.path.exists(path):
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        seen = Counter()
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if isinstance(c, str) and len(c) > 40:
                        for m in NOMBRE_PROPIO.finditer(c):
                            seen[m.group(0)] += 1
        wb.close()
        if seen:
            print("  %s: %d candidatos a nombre propio" % (fname, sum(seen.values())))
            for k, v in seen.most_common(8):
                print("      %dx  %s" % (v, k))


if __name__ == "__main__":
    sys.exit(main())
