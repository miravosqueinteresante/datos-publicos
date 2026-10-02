#!/usr/bin/env python3
"""Consolida los reclamos ciudadanos de la Municipalidad de Asunción (2023-2026).

Lee los 4 .xlsx entregados como respuesta al pedido ID 106387 (Ley 5282/2014)
y produce `www/datos/reclamos.json`, un resumen AGREGADO (sin registros
individuales, sin datos personales) listo para consumir desde Jekyll.

Fuente: Departamento de Atención al Ciudadano — Municipalidad de Asunción.
Uso:
    python scripts/consolidar_reclamos.py --fuente "C:\\...\\Reclamos Municipalidad"
"""

import argparse
import json
import os
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl

REPO_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(REPO_DIR, "www", "datos", "reclamos.json")

# archivo -> (anio, sheet_pendiente, sheet_finalizado, mapping)
FILES = {
    "1790948115_1_AO2023.xlsx": {
        "anio": 2023,
        "pendientes": "PENDIENTES",
        "finalizados": "FINALIZADOS",
        "cols": {
            "pendientes": {"nro": "Nro Reclamo", "tipo": "Tipo de reclamo",
                           "dep": "Dependencia", "fecha": "Fecha de derivacion"},
            "finalizados": {"nro": "Nro Reclamo", "fecha": "Fecha de realización",
                            "tipo": "Tipo de reclamo", "dep": "Dependencia"},
        },
    },
    "1790948115_2_AO2024.xlsx": {
        "anio": 2024,
        "pendientes": "PENDIENTES",
        "finalizados": "FINALIZADOS",
        "cols": {
            "pendientes": {"nro": "Nro Reclamo", "estado": "Estado",
                           "fecha": "Fecha de realización", "tipo": "Tipo de reclamo",
                           "dep": "Dependencia ", "barrio": "Barrio y direccion"},
            "finalizados": {"nro": "Nro Reclamo", "estado": "Estado",
                            "fecha": "Fecha de realización", "tipo": "Tipo de reclamo",
                            "dep": "Dependencia ", "barrio": "Barrio y direccion"},
        },
    },
    "1790948115_3_AO2025.xlsx": {
        "anio": 2025,
        "pendientes": "PENDIENTE",
        "finalizados": "FINALIZADO",
        "cols": {
            "pendientes": {"nro": "Nro Reclamo", "estado": "Estado", "fecha": "Fecha",
                           "tipo": "Tipo de reclamo", "dep": "Departamento", "barrio": "Barrio"},
            "finalizados": {"nro": "Nro Reclamo", "estado": "Estado", "fecha": "Fecha",
                            "tipo": "Tipo de reclamo", "dep": "Departamento", "barrio": "Barrio"},
        },
    },
    "1790948115_4_AO2026.xlsx": {
        "anio": 2026,
        "pendientes": "PENDIENTES",
        "finalizados": "FINALIZADOS",
        "cols": {
            "pendientes": {"nro": "Nro Reclamo", "estado": "Estado", "fecha": "Fecha",
                           "tipo": "Tipo de reclamo", "dep": "Departamento", "barrio": "Barrio"},
            "finalizados": {"nro": "Nro Reclamo", "estado": "Estado", "fecha": "Fecha",
                            "tipo": "Tipo de reclamo", "dep": "Departamento", "barrio": "Barrio"},
        },
    },
}


def unaccent(s):
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def norm_key(s):
    """Clave de agrupación: sin acentos, mayúsculas, espacios colapsados."""
    if s is None:
        return ""
    s = str(s).strip()
    s = unaccent(s)
    s = " ".join(s.split())
    return s.upper()


def parse_fecha(v):
    """Acepta datetime (2025/2026) o texto DD/MM/YYYY (2023/2024)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def barrio_2024(s):
    """Extrae el barrio del campo mezclado 'Barrio y direccion' de 2024.
    Heurística: 'Bo. X' / 'B.O. X' hasta la coma; si no hay, devuelve None
    para no contaminar el ranking de barrios con direcciones."""
    if not s:
        return None
    s = str(s).strip()
    head = s.split(",")[0].strip()
    low = head.lower()
    for prefix in ("bo.", "b.o.", "barrio"):
        if low.startswith(prefix):
            rest = head[len(prefix):].strip()
            return rest if rest else None
    return None


def iter_rows(path, sheet, mapping, anio, estado_por_defecto):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    hdr = [str(h).strip() if h is not None else "" for h in header]

    def idx(name):
        return hdr.index(name) if name in hdr else None

    i = {k: idx(v) for k, v in mapping.items()}
    for r in rows:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        def g(k):
            j = i.get(k)
            return r[j] if j is not None and j < len(r) else None
        nro = g("nro")
        if nro is None:
            continue
        tipo = norm_key(g("tipo"))
        dep = norm_key(g("dep"))
        fecha = parse_fecha(g("fecha"))
        estado = g("estado")
        if estado is None:
            estado = estado_por_defecto
        else:
            estado = norm_key(estado)
        barrio_raw = g("barrio")
        if anio == 2024:
            barrio = norm_key(barrio_2024(barrio_raw)) or None
        else:
            barrio = norm_key(barrio_raw) or None
        yield {"anio": anio, "estado": estado, "fecha": fecha,
               "tipo": tipo, "dep": dep, "barrio": barrio}
    wb.close()


def aggregate(records):
    por_anio = Counter()
    por_mes = Counter()
    por_estado = Counter()
    por_categoria = Counter()
    por_dependencia = Counter()
    por_barrio = Counter()

    for r in records:
        a = r["anio"]
        por_anio[(a, r["estado"])] += 1
        if r["fecha"]:
            por_mes[(r["fecha"][:7], r["estado"])] += 1
        por_estado[r["estado"]] += 1
        if r["tipo"]:
            por_categoria[r["tipo"]] += 1
        if r["dep"]:
            por_dependencia[r["dep"]] += 1
        if r["barrio"]:
            por_barrio[r["barrio"]] += 1

    def build_anio():
        out = []
        for anio in sorted({a for a, _ in por_anio.keys()}):
            d = {"anio": anio, "pendientes": 0, "finalizados": 0, "anulados": 0,
                 "otros": 0, "total": 0}
            for (aa, est), n in por_anio.items():
                if aa != anio:
                    continue
                d["total"] += n
                if est == "PENDIENTE":
                    d["pendientes"] += n
                elif est == "FINALIZADO":
                    d["finalizados"] += n
                elif est == "ANULADO":
                    d["anulados"] += n
                else:
                    d["otros"] += n
            out.append(d)
        return out

    def build_mes():
        out = []
        for mes in sorted({m for m, _ in por_mes.keys()}):
            d = {"mes": mes, "total": 0, "pendientes": 0, "finalizados": 0}
            for (mm, est), n in por_mes.items():
                if mm != mes:
                    continue
                d["total"] += n
                if est == "PENDIENTE":
                    d["pendientes"] += n
                elif est == "FINALIZADO":
                    d["finalizados"] += n
            out.append(d)
        return out

    def build_top(counter, label):
        return [{"nombre": k, "total": v} for k, v in
                sorted(counter.items(), key=lambda x: (-x[1], x[0]))]

    return {
        "por_anio": build_anio(),
        "por_mes": build_mes(),
        "por_estado": [{"estado": k, "total": v} for k, v in
                       sorted(por_estado.items(), key=lambda x: -x[1])],
        "por_categoria": build_top(por_categoria, "categoria"),
        "por_dependencia": build_top(por_dependencia, "dependencia"),
        "por_barrio": build_top(por_barrio, "barrio"),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fuente", default=r"C:\Users\pc\Desktop\Proyectos\Reclamos Municipalidad",
                        help="Carpeta con los .xlsx fuente")
    args = parser.parse_args()

    records = []
    for fname, spec in FILES.items():
        path = os.path.join(args.fuente, fname)
        if not os.path.exists(path):
            print("  WARN: no encontrado %s (se omite)" % path)
            continue
        anio = spec["anio"]
        pend = list(iter_rows(path, spec["pendientes"], spec["cols"]["pendientes"],
                              anio, "PENDIENTE"))
        fin = list(iter_rows(path, spec["finalizados"], spec["cols"]["finalizados"],
                             anio, "FINALIZADO"))
        print("  %s: %d pendientes, %d finalizados" % (fname, len(pend), len(fin)))
        records.extend(pend)
        records.extend(fin)

    resumen = aggregate(records)
    total = len(records)
    finalizados = sum(1 for r in records if r["estado"] == "FINALIZADO")
    pendientes = sum(1 for r in records if r["estado"] == "PENDIENTE")

    cat = resumen["por_categoria"][0] if resumen["por_categoria"] else {"nombre": "", "total": 0}
    dep = resumen["por_dependencia"][0] if resumen["por_dependencia"] else {"nombre": "", "total": 0}

    payload = {
        "_meta": {
            "titulo": "Reclamos ciudadanos — Municipalidad de Asunción",
            "fuente": "Departamento de Atención al Ciudadano, Municipalidad de Asunción",
            "pedido": "Pedido de acceso a la información pública ID 106387 (Ley 5282/2014)",
            "periodo": "01/01/2023 – 11/09/2026",
            "sincronizado": datetime.utcnow().strftime("%Y-%m-%d"),
            "registros": total,
            "nota": ("Datos anonimizados y agregados. No se publican registros "
                     "individuales ni datos personales. Barrios: cobertura parcial "
                     "(sin barrio en 2023; campo mezclado con dirección en 2024)."),
        },
        "kpis": {
            "total": total,
            "pendientes": pendientes,
            "finalizados": finalizados,
            "anulados": sum(1 for r in records if r["estado"] == "ANULADO"),
            "tasa_resolucion": round(finalizados / total * 100, 1) if total else 0,
            "categoria_top": cat["nombre"],
            "categoria_top_total": cat["total"],
            "dependencia_top": dep["nombre"],
            "dependencia_top_total": dep["total"],
        },
        **resumen,
    }

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print("OK: %s (%d registros agregados)" % (OUT_PATH, total))


if __name__ == "__main__":
    sys.exit(main())
